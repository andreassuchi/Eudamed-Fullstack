from conftest import build_workbook
from eudamed_tool.importer import load_registration
from eudamed_tool.models import IssuingEntityCode


def test_import_ok(sample_workbook):
    result = load_registration(sample_workbook)
    assert result.ok, [str(i) for i in result.issues]
    reg = result.registration
    assert len(reg.basic_udis) == 1
    assert len(reg.devices) == 1
    d = reg.devices[0]
    assert d.issuing_entity_code is IssuingEntityCode.GS1
    # fixed portfolio values
    assert d.sterile is False and d.number_of_reuses == -1 and d.base_quantity == 1
    assert d.trade_names[0].trade_name == "AcmeFlow"
    assert d.market_countries[0].country_code == "DE"


def test_import_missing_file(tmp_path):
    result = load_registration(tmp_path / "nope.xlsx")
    assert not result.ok
    assert "file not found" in result.issues[0].message


def test_import_bad_enum_reported_with_location(tmp_path, sample_data):
    sample_data["Devices"][0]["issuing_entity_code"] = "NOT_AN_ENTITY"
    wb = build_workbook(tmp_path / "bad.xlsx", sample_data)
    result = load_registration(wb)
    assert not result.ok
    issue = next(i for i in result.issues if i.column == "issuing_entity_code")
    assert issue.sheet == "Devices" and issue.row == 2


def test_import_orphan_child_row(tmp_path, sample_data):
    sample_data["EMDN"].append({"udi_di": "00000000000000", "emdn_code": "Z9999"})
    wb = build_workbook(tmp_path / "orphan.xlsx", sample_data)
    result = load_registration(wb)
    assert not result.ok
    assert any("no row in Devices sheet" in i.message for i in result.issues)


def test_original_placed_on_market_derived(tmp_path, sample_data):
    sample_data["MarketCountries"].append({"udi_di": "04012345000012", "country_code": "AT"})
    wb = build_workbook(tmp_path / "markets.xlsx", sample_data)
    result = load_registration(wb)
    assert result.ok, [str(i) for i in result.issues]
    markets = {m.country_code: m.original_placed_on_market
               for m in result.registration.devices[0].market_countries}
    assert markets == {"DE": True, "AT": False}


def test_legacy_original_placed_on_market_column_ignored(tmp_path, sample_data):
    from openpyxl import load_workbook as _load
    wb_path = build_workbook(tmp_path / "legacy.xlsx", sample_data)
    wb = _load(wb_path)
    ws = wb["MarketCountries"]
    col = ws.max_column + 1
    ws.cell(row=1, column=col, value="original_placed_on_market")
    ws.cell(row=2, column=col, value="FALSE")  # contradicts the rule; must be ignored
    wb.save(wb_path)
    result = load_registration(wb_path)
    assert result.ok, [str(i) for i in result.issues]
    assert result.registration.devices[0].market_countries[0].original_placed_on_market is True


def test_import_duplicate_udi(tmp_path, sample_data):
    sample_data["Devices"].append(dict(sample_data["Devices"][0]))
    wb = build_workbook(tmp_path / "dup.xlsx", sample_data)
    result = load_registration(wb)
    assert not result.ok
    assert any("duplicate udi_di" in i.message for i in result.issues)
