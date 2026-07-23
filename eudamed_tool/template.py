"""Generate a blank Excel master-data workbook from the shared layout."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .profile import Profile
from .workbook import get_sheets

REQUIRED_FILL = PatternFill("solid", fgColor="FCE4D6")  # light orange = mandatory
HEADER_FONT = Font(bold=True)
DROPDOWN_ROWS = 200  # rows covered by enum dropdown validation


def create_template(path: str | Path, profile: Profile | None = None) -> Path:
    path = Path(path)
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name, columns in get_sheets(profile).items():
        ws = wb.create_sheet(sheet_name)
        for idx, (header, required, allowed) in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=idx, value=header)
            cell.font = HEADER_FONT
            if required:
                cell.fill = REQUIRED_FILL
            ws.column_dimensions[get_column_letter(idx)].width = max(len(header) + 2, 14)
            if allowed:
                dv = DataValidation(
                    type="list",
                    formula1='"' + ",".join(allowed) + '"',
                    allow_blank=not required,
                    showErrorMessage=True,
                )
                col = get_column_letter(idx)
                dv.add(f"{col}2:{col}{DROPDOWN_ROWS + 1}")
                ws.add_data_validation(dv)
        ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    out = create_template(Path(__file__).resolve().parent.parent / "templates" / "eudamed_master_data.xlsx")
    print(f"Template written to {out}")
