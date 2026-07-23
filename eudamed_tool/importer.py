"""Read the Excel master-data workbook into the typed Registration model.

Every problem is reported as an ImportIssue with sheet / row / column context so
the user can fix the workbook cell by cell. Pydantic validation errors are
translated into the same issue format.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook
from pydantic import ValidationError

from .models import BasicUDI, Device, EMDNCode, MarketCountry, ProductionIdentifierType, Registration, TradeName
from .workbook import SHEETS

TRUE_VALUES = {"true", "yes", "y", "1", "x"}
FALSE_VALUES = {"false", "no", "n", "0", ""}


@dataclass
class ImportIssue:
    sheet: str
    row: Optional[int]  # 1-based Excel row; None for sheet-level issues
    column: Optional[str]
    message: str

    def __str__(self) -> str:
        loc = self.sheet
        if self.row is not None:
            loc += f" row {self.row}"
        if self.column:
            loc += f", column '{self.column}'"
        return f"[{loc}] {self.message}"


class ImportResult:
    def __init__(self, registration: Optional[Registration], issues: List[ImportIssue]):
        self.registration = registration
        self.issues = issues

    @property
    def ok(self) -> bool:
        return self.registration is not None and not self.issues


def _cell_to_str(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, str):
        s = value.strip()
        return s if s else None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # Excel stores numbers as floats; avoid "1.0"
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _parse_bool(raw: str, sheet: str, row: int, column: str, issues: List[ImportIssue]) -> Optional[bool]:
    v = raw.strip().lower()
    if v in TRUE_VALUES:
        return True
    if v in FALSE_VALUES:
        return False
    issues.append(ImportIssue(sheet, row, column, f"cannot interpret '{raw}' as TRUE/FALSE"))
    return None


BOOL_COLUMNS = {
    header
    for columns in SHEETS.values()
    for header, _req, allowed in columns
    if allowed == ["TRUE", "FALSE"]
}

# BasicUDI criteria fixed to FALSE for this portfolio; not workbook columns.
# The domain model and XML output keep them (XSD-required), always false.
FIXED_FALSE_FIELDS = {
    "animal_tissues_cells", "human_tissues_cells", "human_product_check",
    "medicinal_product_check", "administering_medicine", "implantable", "reusable",
}

# Device fields fixed for this portfolio; not workbook columns.
# Flags always FALSE; number_of_reuses = -1 (not defined); base_quantity = 1.
FIXED_DEVICE_FIELDS = {
    "sterile", "sterilization", "latex", "reprocessed", "single_use",
    "number_of_reuses", "base_quantity",
}

# Columns that older workbooks may still contain; ignored without error.
# original_placed_on_market is derived (DE -> TRUE, others -> FALSE);
# fixed criteria/values are pinned regardless of what an old workbook says.
LEGACY_IGNORED_COLUMNS = {
    "MarketCountries": {"original_placed_on_market"},
    "BasicUDI": set(FIXED_FALSE_FIELDS),
    "Devices": set(FIXED_DEVICE_FIELDS),
}

# Country whose market entry counts as the original placing on the EU market
ORIGINAL_MARKET_COUNTRY = "DE"


def _read_sheet(ws, sheet_name: str, issues: List[ImportIssue]) -> List[Tuple[int, Dict[str, Any]]]:
    """Return [(excel_row, {header: value})] for non-empty rows, validating headers."""
    expected = [h for h, _r, _a in SHEETS[sheet_name]]
    header_row = [(_cell_to_str(c.value) or "") for c in ws[1]]
    header_row = [h for h in header_row if h]
    ignored = LEGACY_IGNORED_COLUMNS.get(sheet_name, set())
    missing = [h for h in expected if h not in header_row]
    unknown = [h for h in header_row if h not in expected and h not in ignored]
    if missing:
        issues.append(ImportIssue(sheet_name, 1, None, f"missing column(s): {', '.join(missing)}"))
    if unknown:
        issues.append(ImportIssue(sheet_name, 1, None, f"unknown column(s): {', '.join(unknown)}"))
    if missing:
        return []
    col_index = {h: i for i, h in enumerate(header_row)}
    rows: List[Tuple[int, Dict[str, Any]]] = []
    for excel_row, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        record: Dict[str, Any] = {}
        for header in expected:
            idx = col_index[header]
            raw = values[idx] if idx < len(values) else None
            s = _cell_to_str(raw)
            if s is None:
                continue
            if header in BOOL_COLUMNS:
                b = _parse_bool(s, sheet_name, excel_row, header, issues)
                if b is not None:
                    record[header] = b
            else:
                record[header] = s
        if record:
            rows.append((excel_row, record))
    return rows


def _pydantic_issues(err: ValidationError, sheet: str, row: int, issues: List[ImportIssue]) -> None:
    for e in err.errors():
        column = str(e["loc"][0]) if e["loc"] else None
        issues.append(ImportIssue(sheet, row, column, e["msg"]))


def load_registration(path: str | Path) -> ImportResult:
    path = Path(path)
    issues: List[ImportIssue] = []
    if not path.exists():
        return ImportResult(None, [ImportIssue("workbook", None, None, f"file not found: {path}")])
    wb = load_workbook(path, data_only=True)

    missing_sheets = [s for s in SHEETS if s not in wb.sheetnames]
    if missing_sheets:
        issues.append(ImportIssue("workbook", None, None, f"missing sheet(s): {', '.join(missing_sheets)}"))
        return ImportResult(None, issues)

    raw = {name: _read_sheet(wb[name], name, issues) for name in SHEETS}

    basic_udis: List[BasicUDI] = []
    for row, record in raw["BasicUDI"]:
        try:
            basic_udis.append(BasicUDI(**record))
        except ValidationError as err:
            _pydantic_issues(err, "BasicUDI", row, issues)

    # index child rows by udi_di
    def children(sheet: str) -> Dict[str, List[Tuple[int, Dict[str, Any]]]]:
        out: Dict[str, List[Tuple[int, Dict[str, Any]]]] = {}
        for row, record in raw[sheet]:
            udi = record.pop("udi_di", None)
            if not udi:
                issues.append(ImportIssue(sheet, row, "udi_di", "udi_di is required"))
                continue
            out.setdefault(udi, []).append((row, record))
        return out

    trade_names = children("TradeNames")
    emdn = children("EMDN")
    prod_ids = children("ProductionIdentifiers")
    markets = children("MarketCountries")

    devices: List[Device] = []
    seen_udis = set()
    for row, record in raw["Devices"]:
        udi = record.get("udi_di")
        if udi in seen_udis:
            issues.append(ImportIssue("Devices", row, "udi_di", f"duplicate udi_di '{udi}'"))
            continue
        if udi:
            seen_udis.add(udi)
        record["trade_names"] = []
        for crow, crec in trade_names.pop(udi, []):
            try:
                record["trade_names"].append(TradeName(**crec))
            except ValidationError as err:
                _pydantic_issues(err, "TradeNames", crow, issues)
        record["emdn_codes"] = []
        for crow, crec in emdn.pop(udi, []):
            try:
                record["emdn_codes"].append(EMDNCode(**crec))
            except ValidationError as err:
                _pydantic_issues(err, "EMDN", crow, issues)
        record["production_identifiers"] = []
        for crow, crec in prod_ids.pop(udi, []):
            try:
                record["production_identifiers"].append(ProductionIdentifierType(crec["identifier_type"]))
            except (KeyError, ValueError):
                issues.append(ImportIssue("ProductionIdentifiers", crow, "identifier_type",
                                          f"invalid identifier_type '{crec.get('identifier_type')}'"))
        record["market_countries"] = []
        for crow, crec in markets.pop(udi, []):
            # derived, never read from the workbook: DE is the original market
            crec["original_placed_on_market"] = (
                crec.get("country_code") == ORIGINAL_MARKET_COUNTRY
            )
            try:
                record["market_countries"].append(MarketCountry(**crec))
            except ValidationError as err:
                _pydantic_issues(err, "MarketCountries", crow, issues)
        try:
            devices.append(Device(**record))
        except ValidationError as err:
            _pydantic_issues(err, "Devices", row, issues)

    # child rows referencing a udi_di that has no Devices row
    for sheet, orphans in (("TradeNames", trade_names), ("EMDN", emdn),
                           ("ProductionIdentifiers", prod_ids), ("MarketCountries", markets)):
        for udi, rows in orphans.items():
            for crow, _ in rows:
                issues.append(ImportIssue(sheet, crow, "udi_di", f"udi_di '{udi}' has no row in Devices sheet"))

    if issues:
        return ImportResult(None, issues)
    return ImportResult(Registration(basic_udis=basic_udis, devices=devices), issues)
